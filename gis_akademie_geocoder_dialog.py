# -*- coding: utf-8 -*-

import os
import csv
import math

from qgis.PyQt import uic
from qgis.PyQt.QtWidgets import (QDialog, QFileDialog, QMessageBox, QInputDialog,
                                 QGroupBox, QLabel, QPushButton, QCheckBox, QToolButton,
                                 QVBoxLayout, QHBoxLayout,
                                 QTableWidget, QTableWidgetItem, QHeaderView,
                                 QAbstractItemView)
from qgis.PyQt.QtCore import QUrl, QThread, QVariant, QObject, pyqtSignal, Qt
from qgis.PyQt.QtGui import QDesktopServices, QPalette, QColor

from qgis.core import (
    QgsProject, QgsVectorLayer, QgsVectorFileWriter,
    QgsFields, QgsField, QgsFeature, QgsGeometry,
    QgsPointXY, QgsSettings,
    QgsCoordinateReferenceSystem, QgsCoordinateTransform,
    QgsMessageLog, Qgis
)

# Technical details (URLs, file names, tracebacks) go to the QGIS message
# log (View → Panels → Log Messages), not to the dialog's log window — that
# one is meant for end users.
LOG_TAG = 'GIS-Akademie Geocoder'
_MSG_INFO = getattr(getattr(Qgis, 'MessageLevel', Qgis), 'Info')
_TECH_HINT = ('Technische Details für den Support: QGIS-Menü Ansicht → Bedienfelder → '
              f'Protokollmeldungen, Reiter „{LOG_TAG}“.')

from .geocoder_worker import GeocoderWorker
from .local_geocoder import db_exists, set_db_path, all_source_meta, get_plz_coverage
from . import source_downloader as _sd
from .source_downloader import describe, is_automatic
from .source_update_worker import SourceUpdateWorker
from .source_editor_dialog import SourceEditDialog
from . import plz_resolver as _plz
from .geocoder_worker import PLUGIN_VERSION


def _fmt_date(iso):
    """'2026-10-23T16:02:51+02:00' -> '23.10.2026'"""
    if not iso:
        return ''
    try:
        y, m, d = iso[:10].split('-')
        return f'{d}.{m}.{y}'
    except Exception:
        return iso[:10]


def _fmt_int(n):
    return f'{n:,}'.replace(',', '.')


def _qgis_proxies():
    """Proxy from the QGIS network settings (Einstellungen → Optionen →
    Netzwerk) for the plugin's downloads, or None (then the system settings
    apply)."""
    try:
        from qgis.core import QgsNetworkAccessManager
        from qgis.PyQt.QtNetwork import QNetworkProxy
        from urllib.parse import quote
        proxy = QgsNetworkAccessManager.instance().fallbackProxy()
        types = getattr(QNetworkProxy, 'ProxyType', QNetworkProxy)
        if proxy.type() not in (types.HttpProxy, types.HttpCachingProxy) or not proxy.hostName():
            return None
        auth = (f'{quote(proxy.user(), safe="")}:{quote(proxy.password(), safe="")}@'
                if proxy.user() else '')
        url = f'http://{auth}{proxy.hostName()}:{proxy.port() or 8080}'
        return {'http': url, 'https': url}
    except Exception:
        return None


def _apply_proxy():
    _sd.set_proxies(_qgis_proxies())


def _keep_until_finished(thread, worker):
    """A thread that does not stop in time is never terminated (a thread
    killed inside SQLite leaves the database locked and can hang QGIS). The
    cancelled worker finishes on its own; until then thread and worker are kept
    alive in qgis.utils (survives a plugin reload)."""
    import qgis.utils as _qu
    keep = getattr(_qu, '_gis_akademie_geocoder_threads', None)
    if keep is None:
        keep = _qu._gis_akademie_geocoder_threads = []
    keep.append((thread, worker))

    def _release(t=thread):
        for item in list(keep):
            if item[0] is t:
                keep.remove(item)
    thread.finished.connect(_release)

SETTINGS_KEY = 'gis_akademie_geocoder/db_path'
PLZ_SETTINGS_KEY = 'gis_akademie_geocoder/plz_quelle'   # user-defined postcode (PLZ) file (JSON)

# ────────────────────────────────────────────────────────────────────
# Field types, version-safe: QgsField accepts QMetaType.Type only from QGIS
# 3.38 on; older QGIS 3 versions need QVariant.Type (PyQt5 knows both enums,
# so the QGIS version decides, not the import). QGIS 4 / Qt6 removed the
# QVariant.Type enums.
def _field_types():
    if Qgis.QGIS_VERSION_INT >= 33800:
        try:
            from qgis.PyQt.QtCore import QMetaType
            return QMetaType.Type.QString, QMetaType.Type.Int, QMetaType.Type.Double
        except (ImportError, AttributeError):
            pass
    return QVariant.String, QVariant.Int, QVariant.Double


FT_STRING, FT_INT, FT_DOUBLE = _field_types()

FORM_CLASS, _ = uic.loadUiType(
    os.path.join(os.path.dirname(__file__), 'gis_akademie_geocoder_dialog.ui')
)


class CatalogCheckWorker(QObject):
    """Check the online source list in the background (network access)."""
    finished = pyqtSignal(str, str)   # status, text

    def run(self):
        try:
            status, text = _sd.fetch_online_catalog(user_agent_version=PLUGIN_VERSION)
        except Exception as e:
            status, text = 'fehler', f'Die Quellenliste konnte nicht geprüft werden: {e}'
        self.finished.emit(status, text)


class DbStatsWorker(QObject):
    """Computes the DB statistics (addresses per Bundesland) in a
    background thread, so that the UI does not freeze with very large
    databases (millions of addresses)."""
    finished = pyqtSignal(dict, str, int)   # stats, db_path, generation

    def __init__(self, db_path, generation=0):
        super().__init__()
        self._db_path = db_path
        self._generation = generation

    def run(self):
        try:
            from . import local_geocoder as _lc
            stats = _lc.get_db_stats(self._db_path)
        except Exception:
            stats = {}
        self.finished.emit(stats, self._db_path, self._generation)


_METER_IN_DEG = 1.0 / 111320.0
_OFFSET_ANGLES = [i * (2 * math.pi / 8) for i in range(8)]


_WGS84_TRANSFORMS = {}


def _to_wgs84(x, y, epsg):
    """Projected coordinate (x, y) in EPSG:<epsg> to WGS84 (lon, lat) — via
    QgsCoordinateTransform with the project's transform context, created
    once per EPSG code and cached.

    Replaces the former hand-written UTM series expansion, which only
    handled EPSG:25832/25833 correctly and silently returned nonsense for
    any other code (e.g. 4258, 31467).
    Raises ValueError for an unknown/invalid EPSG code."""
    try:
        epsg = int(epsg)
    except (TypeError, ValueError):
        raise ValueError(f'Ungültiger EPSG-Code: {epsg!r}')
    tr = _WGS84_TRANSFORMS.get(epsg)
    if tr is None:
        src = QgsCoordinateReferenceSystem.fromEpsgId(epsg)
        if not src.isValid():
            raise ValueError(f'Unbekanntes Koordinatensystem EPSG:{epsg}')
        dst = QgsCoordinateReferenceSystem.fromEpsgId(4326)
        tr  = QgsCoordinateTransform(src, dst, QgsProject.instance().transformContext())
        _WGS84_TRANSFORMS[epsg] = tr
    pt = tr.transform(QgsPointXY(float(x), float(y)))
    return pt.x(), pt.y()


class GISAkademieGeocoderDialog(QDialog, FORM_CLASS):

    def __init__(self, iface, parent=None):
        super().__init__(parent or iface.mainWindow())
        self.iface = iface
        self.setupUi(self)

        # Both tabs sit in a scroll area. Its viewport paints the lighter
        # "Base" color by default, which makes the background stand out from
        # the rest of the dialog. Instead we fill the viewport with the window
        # color (palette "Window") — adapts to light/dark themes and (unlike a
        # stylesheet) does not recolor any child widgets.
        try:
            _win_role = QPalette.ColorRole.Window   # Qt6 (scoped enums)
        except AttributeError:
            _win_role = QPalette.Window             # Qt5
        # Fill scroll viewports and ALL tab pages (main as well as column-mapping
        # tabs) with the window color. This way white is reserved exclusively for
        # real input/output fields (viewports and tab pages otherwise paint the
        # light "Base" color — among other places behind the bottom button bar).
        for _sa in (self.geocode_scroll, self.admin_scroll):
            try:
                _vp = _sa.viewport()
                _vp.setBackgroundRole(_win_role)
                _vp.setAutoFillBackground(True)
            except Exception:
                pass
        for _w in (self.tab_geocode, self.tab_admin, self.tab_multi, self.tab_single):
            try:
                _w.setBackgroundRole(_win_role)
                _w.setAutoFillBackground(True)
            except Exception:
                pass
        # Render tabs flat (document mode): removes the light frame around the
        # tab pages (sometimes visible on one side only) and the white fill of
        # the active tab — without a stylesheet, theme-compliant.
        for _tw in (self.tab_widget, self.tabs_mapping):
            try:
                _tw.setDocumentMode(True)
            except Exception:
                pass

        self._rows    = []
        self._headers = []
        self._worker  = None
        self._thread  = None
        self._results = []
        self._loading = False

        self._imp_worker = None
        self._imp_thread = None
        self._cancel_requested = False  # geocoding run cancelled by the user
        self._renamed_headers = []     # columns renamed while loading (for the log)

        self._combos = {
            'strasse':    self.cb_strasse,
            'hausnummer': self.cb_hausnummer,
            'plz':        self.cb_plz,
            'stadt':      self.cb_stadt,
            'land':       self.cb_land,
        }

        self._connect_signals()
        self._last_db_stats = None
        self._build_admin_page()
        # Fill the DB path field with the current path
        from . import local_geocoder as _lc
        self.le_db_path.setText(_lc.DB_PATH)
        self._stats_jobs = []          # [(QThread, DbStatsWorker)] — keep references!
        self._stats_generation = 0     # only the most recent result is shown
        self._refresh_db_status()
        self._refresh_source_views()

    # ───────────────────────────────── Lifecycle ──

    def _running_threads(self):
        return [t for t in (self._thread, self._imp_thread)
                if t is not None and t.isRunning()]

    def shutdown(self, wait_ms=5000):
        """Shut down all background threads in a controlled way (closing the
        dialog, unloading the plugin). Workers are cancelled (imports roll back
        their transaction; the database stays unchanged) and get some time to
        stop. A thread that is still busy (e.g. a long SQLite statement) is never
        terminated but kept alive until it ends — terminating it inside SQLite
        would leave the database locked."""
        jobs = [(self._thread, self._worker), (self._imp_thread, self._imp_worker)]
        jobs += list(self._stats_jobs)
        if self._catalog_job is not None:
            jobs.append(self._catalog_job)
        for thread, worker in jobs:
            if worker is not None:
                if hasattr(worker, 'cancel'):
                    try:
                        worker.cancel()
                    except Exception:
                        pass
                # After closing, do not process/write any more results — but let
                # the worker still end its thread.
                for sig in ('finished', 'result_ready', 'progress', 'log', 'detail'):
                    try:
                        getattr(worker, sig).disconnect()
                    except Exception:
                        pass
                if thread is not None:
                    try:
                        worker.finished.connect(thread.quit)
                    except Exception:
                        pass
            if thread is not None and thread.isRunning():
                thread.quit()
                if not thread.wait(wait_ms):
                    _keep_until_finished(thread, worker)
        self._worker = self._thread = None
        self._imp_worker = self._imp_thread = None
        self._stats_jobs = []
        self._catalog_job = None
        self._update_busy_state()

    def closeEvent(self, event):
        if self._running_threads():
            ans = QMessageBox.question(
                self, 'Vorgang läuft',
                'Eine Geocodierung oder ein Import läuft noch.\n'
                'Vorgang abbrechen und Fenster schließen?',
                QMessageBox.StandardButton.Yes | QMessageBox.StandardButton.No,
                QMessageBox.StandardButton.No)
            if ans != QMessageBox.StandardButton.Yes:
                event.ignore()
                return
            self.shutdown()
        super().closeEvent(event)

    # ─────────────────────────────────────────── Signals ──

    def _connect_signals(self):
        # Geocoding tab
        self.btn_csv.clicked.connect(self._browse_input)
        self.btn_gpkg.clicked.connect(self._browse_gpkg)
        self.btn_run.clicked.connect(self._run)
        self.btn_stop.clicked.connect(self._cancel)
        self.btn_help.clicked.connect(self._show_help)
        self.btn_help_admin.clicked.connect(self._show_help)
        self.le_csv.textChanged.connect(self._validate_input_path)
        self.le_layer_name.textChanged.connect(self._validate_layer_name)
        # Tab „Datenbank & Quellen“
        self.btn_db_browse.clicked.connect(self._browse_db_path)
        self.btn_db_save.clicked.connect(self._save_db_path)
        self.btn_import_stop.clicked.connect(self._import_cancel)

    def _show_help(self):
        help_path = os.path.join(os.path.dirname(__file__), 'help.html')
        QDesktopServices.openUrl(QUrl.fromLocalFile(help_path))

    # ───────────────────────────────── DB path ──

    def _browse_db_path(self):
        """A single dialog (as in QGIS): select an existing .db or enter a
        new file name. Overwrite confirmation is disabled, because an existing
        DB is opened (not overwritten) when selected."""
        from . import local_geocoder as _lc
        start_dir = os.path.dirname(_lc.DB_PATH)

        dlg = QFileDialog(self, 'Geocodierungsdatenbank auswählen oder neu anlegen')
        dlg.setAcceptMode(QFileDialog.AcceptMode.AcceptSave)
        dlg.setNameFilter('SQLite-Datenbank (*.db)')
        dlg.setDefaultSuffix('db')
        dlg.setOption(QFileDialog.Option.DontConfirmOverwrite, True)
        if start_dir:
            dlg.setDirectory(start_dir)
        if dlg.exec():
            selected = dlg.selectedFiles()
            if selected:
                path = selected[0]
                if not path.lower().endswith('.db'):
                    path += '.db'
                self.le_db_path.setText(path)

    def _save_db_path(self):
        """Stores the entered DB path in QgsSettings, makes it the active
        path immediately and refreshes the status display."""
        path = self.le_db_path.text().strip()
        if not path:
            self.lbl_db_path_info.setText('⚠ Bitte einen Pfad eingeben.')
            self.lbl_db_path_info.setVisible(True)
            return
        if not path.lower().endswith('.db'):
            path += '.db'
            self.le_db_path.setText(path)

        # Save the path and activate it immediately
        settings = QgsSettings()
        settings.setValue(SETTINGS_KEY, path)
        set_db_path(path)
        self.lbl_db_path_info.setVisible(True)

        exists = os.path.exists(path)
        if exists:
            self.lbl_db_path_info.setText('✓ Vorhandene Datenbank geladen und aktiv.')
        else:
            self.lbl_db_path_info.setText(
                '✓ Pfad gespeichert. Datei wird beim ersten Import angelegt.')

        self._refresh_db_status()
        self._refresh_source_views()

    # ─────────────────────────────────── DB status ──

    def _refresh_db_status(self):
        from . import local_geocoder as _lc
        db_path = _lc.DB_PATH

        if not os.path.exists(db_path):
            self._last_db_stats = {}
            self.lbl_db_status.setText(
                '<span style="color:#a15c00;">Die Datenbank enthält noch keine Adressen – sie wird '
                'beim ersten Import angelegt.</span>')
            self._update_address_column()
            return

        # Give immediate feedback, load statistics in the background — otherwise
        # the UI freezes with large DBs (millions of addresses). Until they are
        # there, nothing stale is used (e.g. as the state list of a geocoding run).
        self._last_db_stats = None
        self._update_address_column()
        self.lbl_db_status.setText(
            '<span style="color:gray;">Lese Statistik … (bei großen Datenbanken kann das einen '
            'Moment dauern)</span>')

        # Do NOT block on a possibly still running stats thread with wait() —
        # that froze the UI for exactly as long as the GROUP BY took. Instead,
        # a generation counter: only the result of the most recent request is
        # shown; older threads run out and clean up after themselves.
        #
        # IMPORTANT: thread AND worker must stay referenced (self._stats_jobs),
        # otherwise Python garbage-collects the worker immediately and the query
        # never finishes („Lese Statistik …“ stays on screen). The result goes to
        # a method of the dialog (QObject) and is therefore executed in the GUI
        # thread — a lambda would run in the background thread.
        self._stats_generation += 1
        thread = QThread()                 # no parent: may outlive the dialog (see shutdown)
        worker = DbStatsWorker(db_path, self._stats_generation)
        worker.moveToThread(thread)
        thread.started.connect(worker.run)
        worker.finished.connect(self._on_db_stats_ready)
        worker.finished.connect(thread.quit)
        thread.finished.connect(self._on_stats_thread_finished)
        self._stats_jobs.append((thread, worker))
        thread.start()

    def _on_stats_thread_finished(self):
        """Release a finished statistics thread together with its worker."""
        thread = self.sender()
        for job in list(self._stats_jobs):
            if job[0] is thread:
                self._stats_jobs.remove(job)
                try:
                    job[1].deleteLater()
                    job[0].deleteLater()
                except Exception:
                    pass

    def _on_db_stats_ready(self, stats, db_path, generation=0):
        # Outdated result (requested again in the meantime) or different path →
        # discard.
        if generation and generation != self._stats_generation:
            return
        from . import local_geocoder as _lc
        if db_path != _lc.DB_PATH:
            return

        self._last_db_stats = stats or {}
        self._update_address_column()
        if not stats:
            self.lbl_db_status.setText(
                '<span style="color:#a15c00;">Die Datenbank enthält noch keine Adressen.</span>')
            return
        total = sum(stats.values())
        try:
            size = f' · Dateigröße {os.path.getsize(db_path) / 1e9:.1f} GB'.replace('.', ',')
        except OSError:
            size = ''
        self.lbl_db_status.setText(
            f'✓ <b>{_fmt_int(total)} Adressen</b> aus {len(stats)} '
            f'{"Bundesland" if len(stats) == 1 else "Bundesländern"}{size}')

    # ─────────────────────────────────── Validation ──

    def _validate_input_path(self, text):
        text = text.strip()
        if not text:
            self._clear_style(self.le_csv)
            return
        valid_exts = ('.csv', '.txt', '.tsv', '.tab', '.xlsx', '.xls', '.ods')
        ext = os.path.splitext(text)[1].lower()
        if not os.path.isfile(text):
            self._set_error(self.le_csv, 'Datei nicht gefunden.')
        elif ext not in valid_exts:
            self._set_error(self.le_csv, f'Format "{ext}" nicht unterstützt.')
        else:
            self._clear_style(self.le_csv)
            if not self._loading:
                self._load_input(text)

    def _validate_layer_name(self, text):
        import re as _re
        text = text.strip()
        if not text:
            self._clear_style(self.le_layer_name)
            return
        errors = []
        if not _re.match(r'^[A-Za-z_][A-Za-z0-9_]*$', text):
            errors.append('Nur Buchstaben, Ziffern und _ erlaubt.')
            errors.append('Keine Leerzeichen. Muss mit Buchstabe oder _ beginnen.')
        if len(text) > 64:
            errors.append('Maximal 64 Zeichen.')
        RESERVED = {'select','from','where','table','index','create','drop',
                    'insert','update','delete','view','trigger','geometry'}
        if text.lower() in RESERVED:
            errors.append(f'"{text}" ist ein reserviertes SQL-Schlüsselwort.')
        if errors:
            self._set_error(self.le_layer_name, '\n'.join(errors))
        else:
            self._clear_style(self.le_layer_name)

    def _set_error(self, w, tip):
        w.setStyleSheet('QLineEdit{border:1px solid #cc0000;background:#fff0f0;}')
        w.setToolTip(tip)

    def _clear_style(self, w):
        w.setStyleSheet('')
        w.setToolTip('')

    # ─────────────────────────────── Load input ──

    def _browse_input(self):
        path, _ = QFileDialog.getOpenFileName(
            self, 'Eingabedatei auswählen', '',
            'Alle unterstützten Formate (*.csv *.txt *.tsv *.tab *.xlsx *.xls *.ods);;'
            'CSV-Dateien (*.csv);;Tab-getrennte Dateien (*.txt *.tsv *.tab);;'
            'Excel-Dateien (*.xlsx *.xls);;OpenDocument-Tabellen (*.ods);;Alle Dateien (*)'
        )
        if path:
            self._loading = True
            self.le_csv.setText(path)
            self._loading = False
            self._load_input(path)

    def _load_input(self, path):
        ext = os.path.splitext(path)[1].lower()
        if ext in ('.xlsx', '.xls'):
            self._load_excel(path)
        elif ext == '.ods':
            self._load_sheet_gdal(path, 'ODS', 'OpenDocument-Tabelle')
        else:
            self._load_text(path)

    # Output fields of the result layer. INPUT columns with the same name are
    # renamed on load (e.g. 'name' -> 'name_in'), otherwise QgsFields.append
    # silently fails and the geocoder value overwrites the user's value.
    _RESERVED_OUTPUT_FIELDS = {
        'fid', 'geom', 'geocoder_address', 'geocoder_source', 'geocoded', 'name',
        'addresstype', 'importance', 'type', 'category', 'accuracy_level',
        'accuracy_label', 'local_score', 'local_score_label', 'lat', 'lon',
    }

    def _sanitize_headers(self, raw_headers):
        """Make column names unique and GeoPackage-compatible: empty names
        become 'spalte_N', duplicates get '_2', '_3' …, collisions with output
        fields get the suffix '_in'. Returns (headers, renamings)."""
        out, seen, renamed = [], set(), []
        for i, h in enumerate(raw_headers):
            orig = '' if h is None else str(h).strip()
            name = orig or f'spalte_{i + 1}'
            base = name
            if name.lower() in self._RESERVED_OUTPUT_FIELDS:
                name = f'{base}_in'
            k = 2
            while name.lower() in seen:
                name = f'{base}_{k}'
                k += 1
            seen.add(name.lower())
            out.append(name)
            if name != orig:
                renamed.append((orig or '(leer)', name))
        return out, renamed

    @staticmethod
    def _cell_to_str(v):
        """Convert a cell value to text. Excel returns numbers as float
        (PLZ 81377 -> '81377.0', house number 7 -> '7.0') and dates as
        datetime — both would make the lookup fail."""
        if v is None:
            return ''
        if isinstance(v, bool):
            return str(v)
        if isinstance(v, float):
            return str(int(v)) if v.is_integer() else repr(v)
        if hasattr(v, 'strftime'):
            try:
                return v.date().isoformat() if hasattr(v, 'date') else v.isoformat()
            except Exception:
                return str(v)
        return str(v).strip()

    def _set_table(self, raw_headers, data_rows, cell_fn=None):
        """Take header row + data rows (sequences) into self._headers/self._rows.
        Short rows are padded with '', surplus fields (without a header) are
        discarded instead of ending up under the key None."""
        cell_fn = cell_fn or self._cell_to_str
        self._headers, self._renamed_headers = self._sanitize_headers(raw_headers)
        n = len(self._headers)
        self._rows = []
        for row in data_rows:
            row = list(row)
            if not any((c is not None and str(c).strip()) for c in row):
                continue   # skip completely empty row
            self._rows.append({self._headers[i]: (cell_fn(row[i]) if i < len(row) else '')
                               for i in range(n)})

    def _load_excel(self, path):
        ext = os.path.splitext(path)[1].lower()
        try:
            if ext == '.xls':
                # Old binary format: only xlrd can read it. (Previously openpyxl
                # was tried, which for .xls raises an InvalidFileException rather
                # than an ImportError — so the xlrd fallback was never reached
                # and .xls was effectively unreadable.)
                try:
                    import xlrd
                except ImportError:
                    QMessageBox.critical(self, 'Fehler',
                        'Für .xls-Dateien wird das Paket "xlrd" benötigt.\n'
                        'Installation in der OSGeo4W-Shell:\n'
                        '    python -m pip install xlrd\n'
                        'Alternativ die Datei in Excel als .xlsx speichern.')
                    return
                wb         = xlrd.open_workbook(path)
                sheetnames = wb.sheet_names()
                if len(sheetnames) > 1:
                    chosen, ok = QInputDialog.getItem(
                        self, 'Tabellenblatt auswählen',
                        f'Die Datei enthält {len(sheetnames)} Blätter. Bitte auswählen:',
                        sheetnames, 0, False)
                    if not ok:
                        return
                else:
                    chosen = sheetnames[0]
                ws   = wb.sheet_by_name(chosen)
                rows = [tuple(ws.row_values(i)) for i in range(ws.nrows)]
            else:
                try:
                    import openpyxl
                except ImportError:
                    # Without openpyxl read the file with GDAL's XLSX driver (part of QGIS)
                    self._load_sheet_gdal(path, 'XLSX', 'Excel')
                    return
                wb         = openpyxl.load_workbook(path, read_only=True, data_only=True)
                sheetnames = wb.sheetnames
                if len(sheetnames) > 1:
                    chosen, ok = QInputDialog.getItem(
                        self, 'Tabellenblatt auswählen',
                        f'Die Datei enthält {len(sheetnames)} Blätter. Bitte auswählen:',
                        sheetnames, 0, False)
                    if not ok:
                        wb.close()
                        return
                else:
                    chosen = sheetnames[0]
                ws   = wb[chosen]
                rows = list(ws.iter_rows(values_only=True))
                wb.close()
            sheet_label = chosen
        except Exception as e:
            QMessageBox.critical(self, 'Fehler', f'Excel-Fehler:\n{e}')
            return

        if not rows:
            QMessageBox.warning(self, 'Leer', 'Das Tabellenblatt enthält keine Daten.')
            return
        self._set_table(rows[0], rows[1:])
        self._after_load(f'Excel — Tabellenblatt: "{sheet_label}"')

    def _read_sheet_gdal(self, path, driver):
        """Read a spreadsheet (ODS, XLSX) with GDAL/OGR, which ships with QGIS.
        All cells as text (keeps leading zeros of postcodes such as '01067');
        the caller treats the first row as header. Returns (sheet, rows), or
        None if the user cancelled the sheet selection."""
        from osgeo import gdal
        ds = gdal.OpenEx(path, gdal.OF_VECTOR | gdal.OF_READONLY, allowed_drivers=[driver],
                         open_options=['HEADERS=DISABLE', 'FIELD_TYPES=STRING'])
        if ds is None:
            raise ValueError('Die Datei kann nicht gelesen werden.')
        names = [ds.GetLayer(i).GetName() for i in range(ds.GetLayerCount())]
        if not names:
            raise ValueError('Die Datei enthält kein Tabellenblatt.')
        if len(names) > 1:
            chosen, ok = QInputDialog.getItem(
                self, 'Tabellenblatt auswählen',
                f'Die Datei enthält {len(names)} Blätter. Bitte auswählen:', names, 0, False)
            if not ok:
                return None
        else:
            chosen = names[0]
        lyr = ds.GetLayerByName(chosen)
        n = lyr.GetLayerDefn().GetFieldCount()
        rows = [tuple(f.GetFieldAsString(i) if f.IsFieldSet(i) else '' for i in range(n))
                for f in lyr]
        ds = None
        return chosen, rows

    def _load_sheet_gdal(self, path, driver, label):
        try:
            res = self._read_sheet_gdal(path, driver)
        except Exception as e:
            QMessageBox.critical(self, 'Fehler', f'{label}-Datei nicht lesbar:\n{e}')
            return
        if res is None:
            return
        sheet, rows = res
        if not rows:
            QMessageBox.warning(self, 'Leer', 'Das Tabellenblatt enthält keine Daten.')
            return
        self._set_table(rows[0], rows[1:])
        self._after_load(f'{label} — Tabellenblatt: "{sheet}"')

    @staticmethod
    def _text_encoding(path):
        """Encoding of a text file, checked over the WHOLE file (not just 8 KB —
        otherwise an umlaut further down ended up as 'Lesefehler'). UTF-16 is
        recognised by its byte order mark (Excel "Unicode-Text"); cp1252 before
        latin-1, because latin-1 never fails."""
        with open(path, 'rb') as f:
            head = f.read(4)
        if head[:2] in (b'\xff\xfe', b'\xfe\xff'):
            return 'utf-16'
        for enc in ('utf-8-sig', 'cp1252', 'latin-1'):
            try:
                with open(path, newline='', encoding=enc) as f:
                    while f.read(1 << 20):
                        pass
                return enc
            except UnicodeDecodeError:
                continue
        return 'latin-1'

    @staticmethod
    def _csv_delimiter(sample):
        """Delimiter of a CSV file. The header line decides: a header without
        any delimiter is a single column — then commas inside the values
        („Hauptstraße 5, 12345 Berlin“) are NOT separators. Otherwise the
        candidate that occurs in the header and splits the data lines into the
        same number of fields wins."""
        lines = [ln for ln in sample.splitlines() if ln.strip()]
        if not lines:
            return ';'
        header, data = lines[0], lines[1:20]
        cands = [d for d in (';', '\t', '|', ',') if header.count(d)]
        if not cands:
            return '\x1f'                       # single column: never split
        def score(d):
            n = header.count(d)
            same = sum(1 for ln in data if ln.count(d) == n)
            return (same, n)
        return max(cands, key=score)

    def _load_text(self, path):
        ext      = os.path.splitext(path)[1].lower()
        try:
            encoding = self._text_encoding(path)
            with open(path, newline='', encoding=encoding) as f:
                sample = f.read(65536)
        except OSError as e:
            QMessageBox.critical(self, 'Fehler', f'Lesefehler:\n{e}')
            return
        if ext in ('.txt', '.tsv', '.tab'):
            delimiter = '\t'
            filetype  = 'Tab-getrennte Textdatei'
        else:
            delimiter = self._csv_delimiter(sample)
            filetype  = 'CSV'
            _enc_name = {'utf-8-sig': 'UTF-8', 'utf-8': 'UTF-8', 'cp1252': 'Windows-1252',
                         'latin-1': 'ISO-8859-1', 'utf-16': 'UTF-16'}.get(encoding, encoding)
            _delim_name = {'\t': 'Tabulator', ';': 'Semikolon', ',': 'Komma',
                           '|': 'senkrechter Strich', '\x1f': 'keins (eine Spalte)'}.get(
                               delimiter, f'„{delimiter}“')
            filetype = f'CSV (Trennzeichen {_delim_name}, Zeichensatz {_enc_name})'
        try:
            with open(path, newline='', encoding=encoding) as f:
                reader = csv.reader(f, delimiter=delimiter)
                header = next(reader, None)
                if header is None:
                    QMessageBox.warning(self, 'Leer', 'Die Datei enthält keine Daten.')
                    return
                self._set_table(header, reader, cell_fn=lambda v: v.strip())
        except Exception as e:
            QMessageBox.critical(self, 'Fehler', f'Lesefehler:\n{e}')
            return
        self._after_load(filetype)

    def _after_load(self, filetype_label):
        self.log.clear()
        self.progress_bar.setValue(0)
        self.lbl_progress.setText('Bereit.')
        self.lbl_filetype.setText(f'Erkannter Dateityp: {filetype_label}')
        self._log(f'Geladen: {len(self._rows)} Zeilen, {len(self._headers)} Spalten  [{filetype_label}]')
        for alt, neu in self._renamed_headers:
            self._log(f'  Spalte „{alt}" → „{neu}" (leer, doppelt oder Name eines Ausgabefelds)')
        # First reset ALL combos, then assign — otherwise combos not yet reset
        # still hold the column indices of the previous file
        for cb in self._combos.values():
            cb.clear()
            cb.addItem('— nicht zugeordnet —')
            cb.addItems(self._headers)
            cb.setEnabled(True)
        for key, cb in self._combos.items():
            self._auto_assign(key, cb)
        # Also fill the address column (tab „Eine Spalte")
        self.cb_address.clear()
        self.cb_address.addItem('— nicht zugeordnet —')
        self.cb_address.addItems(self._headers)
        self.cb_address.setEnabled(True)
        self._auto_assign_address()
        # Select the tab automatically: exactly one column → „Eine Spalte",
        # otherwise „Mehrere Spalten". Reset on every load so that the mode
        # does not get stuck when switching between files.
        if len(self._headers) == 1:
            self.cb_address.setCurrentIndex(1)
            self.tabs_mapping.setCurrentIndex(1)
        else:
            self.tabs_mapping.setCurrentIndex(0)
        if self._rows:
            parts = [f'{k}: {v}' for k, v in list(self._rows[0].items())[:5]]
            self.lbl_preview.setText('Erste Zeile: ' + '  |  '.join(parts))
        stem = os.path.splitext(os.path.basename(self.le_csv.text()))[0]
        self.le_layer_name.setText(stem + '_geocodiert')

    # Hints per field, by priority (best first). Matching in three stages
    # (exact → word start → contained), so that the first header with just any
    # substring does not win: previously 'ort' matched „Ortsteil"/„Sortierung"
    # before „Stadt", 'land' matched „Landkreis" and 'str' matched „District".
    _ASSIGN_HINTS = {
        'strasse':    ['strasse', 'straße', 'street', 'str', 'strassenname', 'straßenname',
                       'adresse', 'address', 'anschrift'],
        'hausnummer': ['hausnummer', 'hausnr', 'haus_nr', 'hnr', 'housenumber',
                       'house_number', 'hn', 'nr'],
        'plz':        ['plz', 'postleitzahl', 'postcode', 'postal_code', 'zip', 'postal'],
        'stadt':      ['stadt', 'ort', 'gemeinde', 'city', 'town', 'ortsname', 'wohnort'],
        'land':       ['bundesland', 'land', 'state', 'bl', 'region'],
    }

    def _auto_assign(self, field_key, combo):
        used = {cb.currentIndex() - 1 for k, cb in self._combos.items()
                if k != field_key and cb.currentIndex() > 0}
        headers = [(i, h.lower().strip()) for i, h in enumerate(self._headers)
                   if i not in used]
        hints = self._ASSIGN_HINTS.get(field_key, [])
        for matcher in (lambda h, hint: h == hint,
                        lambda h, hint: h.startswith(hint),
                        lambda h, hint: len(hint) > 2 and hint in h):
            for hint in hints:
                for i, h in headers:
                    if matcher(h, hint):
                        combo.setCurrentIndex(i + 1)
                        return

    def _auto_assign_address(self):
        # Column names that indicate a complete address in a single column
        hints = ['adresse', 'address', 'anschrift', 'full_address', 'fulladdress']
        for i, h in enumerate(self._headers):
            if any(hint in h.lower() for hint in hints):
                self.cb_address.setCurrentIndex(i + 1)
                return

    # ──────────────────────────────────── Output ──

    def _browse_gpkg(self):
        """Select or create a GeoPackage — same behavior as in QGIS:
        a single dialog in which you either select an existing .gpkg or
        enter a new file name. The overwrite confirmation is disabled,
        because saving only adds/overwrites the layer anyway (an existing
        file is kept)."""
        dlg = QFileDialog(self, 'GeoPackage auswählen oder neu anlegen')
        dlg.setAcceptMode(QFileDialog.AcceptMode.AcceptSave)
        dlg.setNameFilter('GeoPackage (*.gpkg)')
        dlg.setDefaultSuffix('gpkg')
        dlg.setOption(QFileDialog.Option.DontConfirmOverwrite, True)
        if dlg.exec():
            selected = dlg.selectedFiles()
            if selected:
                path = selected[0]
                if not path.lower().endswith('.gpkg'):
                    path += '.gpkg'
                self.le_gpkg.setText(path)

    def _get_col_map(self):
        # Tab index 1 = „Eine Spalte" (combined address), 0 = „Mehrere Spalten"
        if self.tabs_mapping.currentIndex() == 1:
            col = self.cb_address.currentText() if self.cb_address.currentIndex() > 0 else ''
            return {'combined': col}
        return {key: (cb.currentText() if cb.currentIndex() > 0 else '')
                for key, cb in self._combos.items()}

    # ─────────────────────────────── Geocoding ──

    def _run(self):
        if not self._rows:
            QMessageBox.warning(self, 'Eingabe fehlt', 'Bitte zuerst eine Datei laden.')
            return
        col_map = self._get_col_map()
        if 'combined' in col_map:
            if not col_map.get('combined'):
                QMessageBox.warning(self, 'Eingabe fehlt', 'Bitte die Adressspalte zuordnen.')
                return
        elif not col_map.get('strasse') and not col_map.get('stadt'):
            QMessageBox.warning(self, 'Eingabe fehlt', 'Bitte Straße oder Stadt zuordnen.')
            return
        gpkg_path  = self.le_gpkg.text().strip()
        layer_name = self.le_layer_name.text().strip()
        if not gpkg_path:
            QMessageBox.warning(self, 'Eingabe fehlt', 'Bitte GeoPackage-Pfad angeben.')
            return
        if not layer_name:
            QMessageBox.warning(self, 'Eingabe fehlt', 'Bitte Layer-Namen angeben.')
            return

        if self._import_running():
            QMessageBox.information(self, 'Import läuft',
                                    'Während Daten in die Datenbank importiert werden, ist sie '
                                    'gesperrt. Bitte warten, bis die Aktualisierung abgeschlossen ist.')
            return
        self._cancel_requested = False
        self.progress_bar.setValue(0)
        self.progress_bar.setMaximum(len(self._rows))
        self._results = []
        self.log.clear()
        self._log(f'Starte Geocodierung für {len(self._rows)} Adressen …')
        # Use the cached statistics from the admin tab instead of running another
        # GROUP BY scan over millions of rows in the UI thread.
        cached = getattr(self, '_last_db_stats', None)
        local_only = self.cb_local_only.isChecked()

        if cached:
            self._log(f'Lokale Datenbank mit {len(cached)} Bundesländern.')
        elif db_exists():
            self._log('Lokale Datenbank vorhanden.')
        elif local_only:
            self._log('⚠ DSGVO-Modus aktiv, aber keine lokale Datenbank vorhanden — '
                      'es können keine Adressen gefunden werden. Bitte im Reiter '
                      '„Datenbank & Quellen“ Daten laden.')
        else:
            self._log('Keine lokale Datenbank — alle Adressen werden über Nominatim '
                      '(OpenStreetMap) gesucht.')

        if local_only:
            self._log('🔒 DSGVO-konformes Geocodieren: nur die lokale Datenbank, '
                      'keine Online-Abfrage.')

        _apply_proxy()
        self._thread = QThread()
        self._worker = GeocoderWorker(
            self._rows, col_map, local_only=local_only,
            bundeslaender=list(cached.keys()) if cached else None)
        self._worker.moveToThread(self._thread)
        self._thread.started.connect(self._worker.run)
        self._worker.progress.connect(self._on_progress)
        self._worker.result_ready.connect(self._on_results)
        self._worker.detail.connect(self._tech_log)
        self._worker.finished.connect(self._on_finished)
        self._thread.start()
        self._update_busy_state(geocoding_running=True)

    def _cancel(self):
        if self._worker:
            self._worker.cancel()
        self._cancel_requested = True
        self._log('Abbruch angefordert …')
        self.btn_stop.setEnabled(False)

    def _on_progress(self, current, total, address):
        self.progress_bar.setValue(current)
        self.lbl_progress.setText(f'Geocodiere {current} / {total}')
        self._log(f'[{current}/{total}] {address}')

    def _on_results(self, results):
        self._results = results

    def _on_finished(self):
        self._thread.quit()
        self._thread.wait()
        self._update_busy_state(geocoding_running=False)
        abort = getattr(self._worker, 'abort_reason', '') if self._worker is not None else ''
        if abort:
            # e.g. the database is locked by an import running in another QGIS
            self._log(f'⚠ {abort}')
            QMessageBox.warning(self, 'Geocodierung beendet', abort)
        if not self._results:
            self._log('Keine Ergebnisse.')
            return
        ok     = sum(1 for r in self._results if r['geocoded'])
        failed = [r for r in self._results if not r['geocoded']]
        local  = sum(1 for r in self._results if r.get('source','').startswith('Lokal') and r['geocoded'])
        nom    = sum(1 for r in self._results if r.get('source','') == 'Nominatim' and r['geocoded'])
        fallbk = sum(1 for r in self._results if r.get('source','') == 'Nominatim (Fallback)' and r['geocoded'])
        details = f'Lokal: {local}, Nominatim: {nom}'
        if fallbk:
            details += f', Nominatim (Fallback): {fallbk}'
        self._log(f'Fertig: {ok}/{len(self._results)} gefunden  ({details})')

        # Output the reasons for all addresses NOT found locally, grouped —
        # consistently in both modes (with/without Nominatim). In online mode
        # these are the addresses resolved via fallback or not at all; in GDPR
        # (DSGVO) mode all without a local hit. Helps to understand why
        # (Bundesland not imported, municipality missing, street not found in
        # the municipality).
        local_miss = [r for r in self._results
                      if not (r.get('source', '').startswith('Lokal') and r['geocoded'])]
        if local_miss:
            from collections import Counter
            gruende = Counter((r.get('local_diagnose') or 'unbekannt') for r in local_miss)
            self._log(f'Lokal nicht gefunden — Gründe ({len(local_miss)} Adressen):')
            for grund, anzahl in gruende.most_common():
                self._log(f'  • {anzahl}×  {grund}')

        if failed:
            self._log(f'Nicht gefunden ({len(failed)}):')
            for r in failed:
                self._log(f'  ✗  {r["address_query"]}')
        gpkg_path  = self.le_gpkg.text().strip()
        layer_name = self.le_layer_name.text().strip()
        split      = self.cb_split_unfound.isChecked()
        if (self._cancel_requested or abort) and len(self._results) < len(self._rows):
            # A cancelled run never overwrites a (complete) layer of the same name
            layer_name = self._free_layer_name(gpkg_path, f'{layer_name}_teil')
            ans = QMessageBox.question(
                self, 'Geocodierung nicht vollständig',
                f'Die Geocodierung wurde nach {len(self._results)} von {len(self._rows)} Adressen '
                f'beendet.\nTeilergebnis als Layer „{layer_name}“ speichern?',
                QMessageBox.StandardButton.Yes | QMessageBox.StandardButton.No,
                QMessageBox.StandardButton.Yes)
            if ans != QMessageBox.StandardButton.Yes:
                self._log('Abgebrochen – kein Ergebnis gespeichert.')
                return
        try:
            if split:
                found    = [r for r in self._results if r['geocoded']]
                notfound = [r for r in self._results if not r['geocoded']]
                if found:
                    self._save_to_gpkg(found, gpkg_path, layer_name)
                    self._log(f'Gespeichert ({len(found)} Treffer): {gpkg_path}  →  Layer "{layer_name}"')
                    if self.cb_load.isChecked():
                        self._load_to_qgis(gpkg_path, layer_name)
                else:
                    self._log('Keine Treffer — es wurde kein Ergebnis-Layer geschrieben.')
                if notfound:
                    # After the layer: a CSV that is open in Excel must not cost the result
                    try:
                        csv_path = self._write_unfound_csv(notfound, gpkg_path, layer_name)
                        self._log(f'{len(notfound)} nicht gefundene Adresse(n) abgelegt: {csv_path}')
                    except OSError as e:
                        self._log(f'⚠ Die Liste der nicht gefundenen Adressen konnte nicht geschrieben '
                                  f'werden (Datei geöffnet?): {e}')
                else:
                    self._log('Alle Adressen gefunden — keine CSV nötig.')
            else:
                self._save_to_gpkg(self._results, gpkg_path, layer_name)
                self._log(f'Gespeichert: {gpkg_path}  →  Layer "{layer_name}"')
                if self.cb_load.isChecked():
                    self._load_to_qgis(gpkg_path, layer_name)
        except Exception as e:
            self._log(f'Fehler beim Speichern des Ergebnisses: {e}')
            self._tech_log(f'Speichern {gpkg_path} / {layer_name}: {e}')
            QMessageBox.critical(self, 'Fehler', f'Das Ergebnis konnte nicht gespeichert werden:\n{e}')

    def _write_unfound_csv(self, notfound, gpkg_path, layer_name):
        """Writes the addresses that were not found as a CSV next to the
        GeoPackage. File name: <layer name>_nicht_gefunden.csv (semicolon-
        separated, UTF-8 with BOM so that umlauts display correctly in Excel).
        Contains the original input columns plus the assembled query and the
        reason."""
        out_dir  = os.path.dirname(gpkg_path) or '.'
        base     = layer_name.strip() or os.path.splitext(os.path.basename(gpkg_path))[0]
        csv_path = os.path.join(out_dir, f'{base}_nicht_gefunden.csv')
        orig_cols  = list(notfound[0]['original_row'].keys())
        fieldnames = orig_cols + ['abfrage', 'hinweis']
        with open(csv_path, 'w', newline='', encoding='utf-8-sig') as f:
            w = csv.writer(f, delimiter=';')
            w.writerow(fieldnames)
            for r in notfound:
                row     = [r['original_row'].get(c, '') for c in orig_cols]
                hinweis = r.get('local_diagnose') or r.get('accuracy_label', '') or 'nicht gefunden'
                row    += [r.get('address_query', ''), hinweis]
                w.writerow(row)
        return csv_path

    def _free_layer_name(self, gpkg_path, name):
        """name, or name_2, name_3 … if the GeoPackage already has such a layer."""
        existing = {n.lower() for n in self._gpkg_sublayers(gpkg_path)} if os.path.exists(gpkg_path) else set()
        cand, k = name, 2
        while cand.lower() in existing:
            cand = f'{name}_{k}'
            k += 1
        return cand

    # ──────────────────────────────── GeoPackage ──

    def _save_to_gpkg(self, results, gpkg_path, layer_name):
        fields = QgsFields()
        for col in results[0]['original_row'].keys():
            fields.append(QgsField(col, FT_STRING))
        for fname, ftype in [
            ('geocoder_address',   FT_STRING),
            ('geocoder_source',    FT_STRING),
            ('geocoded',           FT_INT),
            ('name',               FT_STRING),
            ('addresstype',        FT_STRING),
            ('importance',         FT_DOUBLE),
            ('type',               FT_STRING),
            ('category',           FT_STRING),
            ('accuracy_level',     FT_INT),
            ('accuracy_label',     FT_STRING),
            ('local_score',        FT_DOUBLE),
            ('local_score_label',  FT_STRING),
            ('lat',                FT_DOUBLE),
            ('lon',                FT_DOUBLE),
        ]:
            fields.append(QgsField(fname, ftype))

        options = QgsVectorFileWriter.SaveVectorOptions()
        options.driverName   = 'GPKG'
        options.fileEncoding = 'UTF-8'
        options.layerName    = layer_name

        # Choose ActionOnExistingFile correctly:
        # - file does not exist  → CreateOrOverwriteFile (creates a new GPKG)
        # - file exists          → CreateOrOverwriteLayer (adds the layer /
        #   replaces only the layer of the same name, other layers are kept)
        # Depending on the QGIS version, the enum is nested under
        # ActionOnExistingFile or sits directly on the class.
        try:
            _Action = QgsVectorFileWriter.ActionOnExistingFile
        except AttributeError:
            _Action = QgsVectorFileWriter
        if os.path.exists(gpkg_path):
            options.actionOnExistingFile = _Action.CreateOrOverwriteLayer
        else:
            options.actionOnExistingFile = _Action.CreateOrOverwriteFile

        mem_layer = self._build_memory_layer(fields, results)
        ctx       = QgsProject.instance().transformContext()
        if hasattr(QgsVectorFileWriter, 'writeAsVectorFormatV3'):      # QGIS ≥ 3.20
            res = QgsVectorFileWriter.writeAsVectorFormatV3(mem_layer, gpkg_path, ctx, options)
        else:                                                           # QGIS 3.16–3.18
            res = QgsVectorFileWriter.writeAsVectorFormatV2(mem_layer, gpkg_path, ctx, options)
        writer, err = res[0], res[1]
        # Depending on the version/Qt, the error enum is scoped
        # (WriterError.NoError) or sits directly on the class.
        _no_error = getattr(getattr(QgsVectorFileWriter, 'WriterError', QgsVectorFileWriter),
                            'NoError')
        if writer != _no_error:
            raise Exception(f'Schreibfehler: {err}')

    def _build_memory_layer(self, fields, results):
        do_offset = self.cb_offset.isChecked()
        layer     = QgsVectorLayer('Point?crs=EPSG:4326', 'tmp', 'memory')
        pr        = layer.dataProvider()
        pr.addAttributes(fields)
        layer.updateFields()

        coord_counter = {}
        feats = []
        n_transform_err = 0
        for r in results:
            feat = QgsFeature()
            feat.setFields(layer.fields())

            lat = lon = None
            if r['geocoded']:
                if r.get('utm_x') is not None:
                    # Local coordinate (UTM 32/33 or other) → WGS84 via QGIS
                    epsg = r.get('epsg') or 25832
                    try:
                        lon, lat = _to_wgs84(r['utm_x'], r['utm_y'], epsg)
                    except Exception as e:
                        n_transform_err += 1
                        if n_transform_err <= 3:
                            self._log('Koordinate konnte nicht umgerechnet werden: '
                                      f'{r.get("address_query", "")}')
                        self._tech_log(f'Transformation EPSG:{epsg} → 4326 fehlgeschlagen: {e}')
                        lon = lat = None
                elif r.get('lon') is not None and r.get('lat') is not None:
                    lon = float(r['lon'])
                    lat = float(r['lat'])

                if lon is not None and lat is not None:
                    # Apply the offset only to the GEOMETRY — the lat/lon fields
                    # keep the original coordinate (as documented).
                    geo_lon, geo_lat = lon, lat
                    if do_offset:
                        key = (round(lat, 7), round(lon, 7))
                        n   = coord_counter.get(key, 0)
                        coord_counter[key] = n + 1
                        if n > 0:
                            angle    = _OFFSET_ANGLES[(n-1) % len(_OFFSET_ANGLES)]
                            geo_lon += (_METER_IN_DEG / math.cos(math.radians(lat))) * math.cos(angle)
                            geo_lat += _METER_IN_DEG * math.sin(angle)
                    feat.setGeometry(QgsGeometry.fromPointXY(QgsPointXY(geo_lon, geo_lat)))

            for col, val in r['original_row'].items():
                feat[col] = val

            feat['geocoder_address']  = r['address_query']
            feat['geocoder_source']   = r.get('source', '')
            feat['geocoded']          = 1 if r['geocoded'] else 0
            feat['name']              = r.get('name', '')
            feat['addresstype']       = r.get('addresstype', '')
            feat['importance']        = r.get('importance')
            feat['type']              = r.get('type', '')
            feat['category']          = r.get('category', '')
            feat['accuracy_level']    = r.get('accuracy_level', 0)
            feat['accuracy_label']    = r.get('accuracy_label', '')
            ls = r.get('local_score')
            feat['local_score']       = round(float(ls), 4) if ls is not None else None
            feat['local_score_label'] = r.get('local_score_label', '')
            feat['lat']               = lat
            feat['lon']               = lon
            feats.append(feat)

        pr.addFeatures(feats)
        return layer

    def _load_to_qgis(self, gpkg_path, layer_name):
        uri   = f'{gpkg_path}|layername={layer_name}'
        layer = QgsVectorLayer(uri, layer_name, 'ogr')
        if layer.isValid():
            QgsProject.instance().addMapLayer(layer)
            self._log(f'Layer "{layer_name}" geladen.')
        else:
            self._log('Warnung: Layer konnte nicht geladen werden.')

    # ─────────────────────────── Automatic source download ──

    @staticmethod
    def _catalog():
        try:
            return _sd.load_catalog()
        except Exception as e:
            return {'sources': {}, '_fehler': [f'Quellenkatalog nicht lesbar: {e}']}

    def _build_admin_page(self):
        """Tab „Datenbank & Quellen“ (a single page):
        geocoding database (path + status) · sources per Bundesland (table
        with an edit button per row, „Alle eingetragenen Quellen
        aktualisieren“) · postcode (PLZ) areas · progress."""
        lay = self.admin_scroll_content.layout()
        idx = lay.indexOf(self.grp_import_progress)
        self.lbl_db_path_info.setVisible(False)     # only after „Pfad speichern“

        # ── Sources per Bundesland ────────────────────────────────────────
        grp = QGroupBox('Quellen je Bundesland')
        g = QVBoxLayout(grp)
        g.setSpacing(6)
        self.tbl_sources = QTableWidget(0, 6)
        self.tbl_sources.setHorizontalHeaderLabels(
            ['Bundesland', 'Adressen', 'Quelle', 'Bezug', 'Zuletzt geladen', ''])
        self.tbl_sources.verticalHeader().setVisible(False)
        self.tbl_sources.verticalHeader().setDefaultSectionSize(28)
        self.tbl_sources.setSelectionBehavior(QAbstractItemView.SelectionBehavior.SelectRows)
        self.tbl_sources.setSelectionMode(QAbstractItemView.SelectionMode.SingleSelection)
        self.tbl_sources.setEditTriggers(QAbstractItemView.EditTrigger.NoEditTriggers)
        self.tbl_sources.setWordWrap(False)
        self.tbl_sources.setVerticalScrollBarPolicy(Qt.ScrollBarPolicy.ScrollBarAlwaysOff)
        self.tbl_sources.setHorizontalScrollBarPolicy(Qt.ScrollBarPolicy.ScrollBarAlwaysOff)
        hh = self.tbl_sources.horizontalHeader()
        for c in range(6):
            hh.setSectionResizeMode(c, QHeaderView.ResizeMode.ResizeToContents)
        hh.setSectionResizeMode(2, QHeaderView.ResizeMode.Stretch)
        g.addWidget(self.tbl_sources)

        row = QHBoxLayout()
        self.cb_auto_force = QCheckBox('Auch ohne neue Daten erneut laden')
        self.cb_auto_force.setToolTip(
            'Normalerweise wird je Quelle nur geladen und importiert, wenn sich der Stand seit dem '
            'letzten Abruf geändert hat (Internet: neuer Stand beim Anbieter; Datei: die Datei '
            'wurde ersetzt). Mit Haken werden alle eingetragenen Quellen neu geladen.')
        row.addWidget(self.cb_auto_force)
        row.addStretch(1)
        self.btn_auto_update = QPushButton('Alle eingetragenen Quellen aktualisieren')
        self.btn_auto_update.setToolTip(
            'Lädt für jedes Bundesland mit eingetragener Quelle die aktuellen Daten und ersetzt '
            'dessen Bestand — nur wenn es neue Daten gibt. Fehlende Postleitzahlen werden dabei '
            'automatisch ergänzt.')
        self.btn_catalog = QPushButton('Quellenliste prüfen')
        self.btn_catalog.setToolTip(
            'Lädt die aktuelle Quellenliste aus dem Repository des Plugins (GitHub). Ändert ein '
            'Land seinen Download-Link, kommt der neue Link so ohne neue Plugin-Version an. '
            'Geschieht auch automatisch vor „Alle eingetragenen Quellen aktualisieren“. '
            'Eigene Quellen bleiben unberührt.')
        row.addWidget(self.btn_catalog)
        row.addWidget(self.btn_auto_update)
        g.addLayout(row)

        self.lbl_auto_info = QLabel('')
        self.lbl_auto_info.setWordWrap(True)
        self.lbl_auto_info.setTextFormat(Qt.TextFormat.RichText)
        g.addWidget(self.lbl_auto_info)

        self.lbl_sources_note = QLabel('')
        self.lbl_sources_note.setWordWrap(True)
        self.lbl_sources_note.setTextFormat(Qt.TextFormat.RichText)
        self.lbl_sources_note.setStyleSheet('color: gray; font-size: 10px;')
        g.addWidget(self.lbl_sources_note)
        lay.insertWidget(idx, grp)
        self.grp_sources = grp

        # ── Postcode (PLZ) areas: bundled (OSM) or user file ───────────────
        grp2 = QGroupBox('Postleitzahlen')
        v2 = QVBoxLayout(grp2)
        v2.setSpacing(6)
        self.lbl_plz_status = QLabel('')
        self.lbl_plz_status.setWordWrap(True)
        self.lbl_plz_status.setTextFormat(Qt.TextFormat.RichText)
        v2.addWidget(self.lbl_plz_status)
        self.lbl_plz_source = QLabel('')
        self.lbl_plz_source.setWordWrap(True)
        self.lbl_plz_source.setTextFormat(Qt.TextFormat.RichText)
        self.lbl_plz_source.setOpenExternalLinks(True)
        v2.addWidget(self.lbl_plz_source)
        row2 = QHBoxLayout()
        self.btn_plz_own = QPushButton('Eigene PLZ-Gebiete …')
        self.btn_plz_own.setToolTip('PLZ-Gebiete aus einer eigenen Datei verwenden (GeoPackage, '
                                    'Shapefile, GeoJSON), z.B. lizenzierte Daten der Deutschen Post.')
        self.btn_plz_bundled = QPushButton('Mitgelieferte verwenden')
        row2.addWidget(self.btn_plz_own)
        row2.addWidget(self.btn_plz_bundled)
        row2.addStretch(1)
        v2.addLayout(row2)
        hint2 = QLabel('Fehlende Postleitzahlen werden beim Import und beim Aktualisieren der '
                       'Quellen automatisch aus den PLZ-Gebieten ergänzt. PLZ aus den amtlichen '
                       'Quellen haben immer Vorrang. Nach einem Wechsel der PLZ-Gebiete werden die '
                       'ergänzten PLZ beim nächsten Aktualisieren neu berechnet.')
        hint2.setWordWrap(True)
        hint2.setStyleSheet('color: gray; font-size: 10px;')
        v2.addWidget(hint2)
        lay.insertWidget(idx + 1, grp2)
        self.grp_plz = grp2

        self.btn_auto_update.clicked.connect(self._auto_update_all)
        self.btn_catalog.clicked.connect(self._check_catalog)
        self._catalog_job = None
        self.tbl_sources.itemDoubleClicked.connect(self._on_source_double_clicked)
        self.btn_plz_own.clicked.connect(self._choose_own_plz)
        self.btn_plz_bundled.clicked.connect(self._use_bundled_plz)

    @staticmethod
    def _edit_icon():
        """Pencil icon from the QGIS theme (editing mode)."""
        try:
            from qgis.core import QgsApplication
            ic = QgsApplication.getThemeIcon('/mActionToggleEditing.svg')
            if ic is not None and not ic.isNull():
                return ic
        except Exception:
            pass
        return None

    @staticmethod
    def _breakable(text):
        """Make long URLs/paths wrappable (zero-width spaces after slash,
        backslash, underscore, ? and &) — otherwise the longest "word"
        determines the minimum width of the whole tab."""
        import html
        t = html.escape(text or '')
        for ch in ('/', '\\', '_', '?', '&amp;'):
            t = t.replace(ch, ch + '&#8203;')
        return t

    # ------------------------------------------------------------------ #
    def _refresh_source_views(self):
        """Refresh the sources table, summary and postcode (PLZ) info. Reads
        only small auxiliary tables — and not during a running import (which
        may hold the write lock)."""
        if self._imp_thread is not None and self._imp_thread.isRunning():
            return
        cat  = self._catalog()
        meta = all_source_meta()
        self._fill_sources_table(cat, meta)
        self._update_sources_summary(cat)
        self._update_plz_status()
        self._update_plz_source_info()
        self._update_busy_state()

    def _update_sources_summary(self, cat):
        reg     = _sd.registered_bundeslaender(cat)
        missing = [bl for bl in _sd.BUNDESLAENDER if bl not in reg]
        txt = (f'Eingetragen: <b>{len(reg)} von {len(_sd.BUNDESLAENDER)}</b> Bundesländern. '
               '<span style="color:gray;">Aktualisiert wird nur, wo es neue Daten gibt.</span>')
        if missing:
            txt += (f'<br><span style="color:#a15c00;">Noch ohne Quelle: {", ".join(missing)}.</span> '
                    'Die Daten gibt es dort nur über ein Portal (Shop, Anmeldung): Datei herunterladen '
                    'und über den Stift eintragen — den Link zum Portal zeigt die Detailseite.')
        if cat.get('_quelle') == 'online':
            liste = (f'Quellenliste: Stand {_fmt_date(cat.get("stand"))}, online nachgeladen'
                     + (f' am {_fmt_date(cat.get("_online_abgerufen"))}' if cat.get('_online_abgerufen') else '')
                     + '.')
        else:
            liste = f'Quellenliste: Stand {_fmt_date(cat.get("stand"))} (mit dem Plugin geliefert).'
        txt += f'<br><span style="color:gray;">{liste}</span>'
        self.lbl_auto_info.setText(txt)

    def _update_plz_status(self):
        """One line: do all addresses have a postcode (PLZ), or where are some still missing?"""
        from . import local_geocoder as _lc
        src = _plz.active_source()
        cov = get_plz_coverage() if os.path.exists(_lc.DB_PATH) else None
        todo = [(bl, ohne) for bl, (n, ohne) in sorted((cov or {}).items()) if n and ohne / n >= 0.001]
        if todo and src:
            txt = ('<span style="color:#a15c00;">Noch ohne PLZ: '
                   + ', '.join(f'{bl} ({_fmt_int(o)} Adressen)' for bl, o in todo)
                   + ' – wird beim nächsten „Alle eingetragenen Quellen aktualisieren“ ergänzt.</span>')
        elif todo:
            txt = ('<span style="color:#a15c00;">Ohne PLZ: '
                   + ', '.join(f'{bl} ({_fmt_int(o)})' for bl, o in todo)
                   + ' – keine PLZ-Gebiete verfügbar.</span>')
        elif cov:
            txt = '✓ Alle Adressen haben eine Postleitzahl.'
        else:
            txt = ''
        self.lbl_plz_status.setText(txt)
        self.lbl_plz_status.setVisible(bool(txt))

    @staticmethod
    def _bezug(e):
        """(text, color) of the column „Bezug“ (source type)."""
        if _sd.is_file_source(e):
            return 'Datei', '#1565c0'
        if is_automatic(e):
            return ('Internet' + (' (experimentell)' if e.get('status') == 'experimentell' else ''),
                    '#2e7d32')
        return 'nicht eingetragen', '#a15c00'

    def _fill_sources_table(self, cat, meta):
        sel = self._selected_source_bl()
        self.tbl_sources.setRowCount(0)
        icon = self._edit_icon()
        for bl in _sd.BUNDESLAENDER:
            e = cat['sources'].get(bl) or {}
            r = self.tbl_sources.rowCount()
            self.tbl_sources.insertRow(r)
            bezug, color = self._bezug(e)
            if _sd.is_file_source(e):
                quelle = os.path.basename(e.get('pfad', ''))
            else:
                quelle = e.get('quelle_name', '—')
                if e.get('_herkunft') == 'eigene':
                    quelle += ' · eigene Quelle'
            m = meta.get(bl)
            zuletzt = _fmt_date(m.get('importiert_am')) if m else ''
            vals = [bl, '', quelle, bezug, zuletzt]
            tip = describe(e) + (f'\n{e.get("url")}' if e.get('url') and not _sd.is_file_source(e) else '')
            for c, val in enumerate(vals):
                it = QTableWidgetItem(val)
                it.setToolTip(tip)
                if c == 3:
                    it.setForeground(QColor(color))
                if c == 1:
                    it.setTextAlignment(int(Qt.AlignmentFlag.AlignRight | Qt.AlignmentFlag.AlignVCenter))
                self.tbl_sources.setItem(r, c, it)
            btn = QToolButton()
            if icon:
                btn.setIcon(icon)
                btn.setToolButtonStyle(Qt.ToolButtonStyle.ToolButtonIconOnly)
            else:
                btn.setText('✎')
            btn.setAutoRaise(True)
            btn.setToolTip(f'{bl} bearbeiten: Quelle eintragen oder ändern, einzeln aktualisieren, '
                           'Datei importieren (auch per Doppelklick auf die Zeile)')
            btn.setProperty('bundesland', bl)
            btn.clicked.connect(self._on_edit_button)
            self.tbl_sources.setCellWidget(r, 5, btn)
        self._update_address_column()
        # Show the whole table without its own scroll bar (the page scrolls)
        hdr = self.tbl_sources.horizontalHeader().sizeHint().height()
        rows = sum(self.tbl_sources.rowHeight(r) for r in range(self.tbl_sources.rowCount()))
        self.tbl_sources.setFixedHeight(hdr + rows + 2 * self.tbl_sources.frameWidth() + 2)
        fehler = cat.get('_fehler') or []
        note = (f'Eigene Quellen werden gespeichert in: {self._breakable(_sd.USER_CATALOG_PATH or "—")} '
                '(übersteht Plugin-Updates). Eine eigene Quelle ersetzt die Standardquelle des '
                'Bundeslands.')
        if fehler:
            note += '<br><span style="color:#c62828;">' + '<br>'.join(fehler) + '</span>'
        self.lbl_sources_note.setText(note)
        if sel:
            for r in range(self.tbl_sources.rowCount()):
                if self.tbl_sources.item(r, 0).text() == sel:
                    self.tbl_sources.selectRow(r)
                    break

    def _update_address_column(self):
        """Column „Adressen“ (addresses) from the most recently read statistics."""
        if not hasattr(self, 'tbl_sources'):
            return
        stats = self._last_db_stats
        for r in range(self.tbl_sources.rowCount()):
            it_bl, it_n = self.tbl_sources.item(r, 0), self.tbl_sources.item(r, 1)
            if it_bl is None or it_n is None:
                continue
            if stats is None:
                it_n.setText('…')
            else:
                n = stats.get(it_bl.text())
                it_n.setText(_fmt_int(n) if n else '–')

    # ------------------------------------------------------------------ #
    def _update_plz_source_info(self):
        src = _plz.active_source()
        if not src:
            self.lbl_plz_source.setText('<span style="color:#a15c00;">Keine PLZ-Gebiete verfügbar – '
                                        'fehlende PLZ werden nicht ergänzt.</span>')
            self.btn_plz_bundled.setEnabled(_plz.get_override() is not None)
            return
        if src['eigene']:
            ov = _plz.get_override() or {}
            text = (f'<b>Eigene PLZ-Datei:</b> {self._breakable(src["path"])}<br>'
                    f'Spalte „{src["field"]}“' + (f', Layer „{ov.get("layer")}“' if ov.get('layer') else '')
                    + f' · Dateistand {_fmt_date(src["stand"])}')
        else:
            m = _plz.bundled_meta()
            text = (f'<b>Mitgeliefert:</b> {src["name"]} · {_fmt_int(m.get("anzahl", 0))} Gebiete<br>'
                    f'Quelle: <a href="{m.get("quelle_url", "")}">{m.get("quelle", "")}</a>'
                    + (f' (<a href="{m.get("auszug_url", "")}">{m.get("auszug")}</a>)'
                       if m.get('auszug') else '')
                    + f' · {m.get("attribution", "")} · Lizenz '
                    f'<a href="{m.get("lizenz_url", "")}">{m.get("lizenz", "")}</a>')
        self.lbl_plz_source.setText(text)
        self.btn_plz_bundled.setEnabled(bool(src['eigene']))

    def _choose_own_plz(self):
        path, _ = QFileDialog.getOpenFileName(
            self, 'Eigene PLZ-Gebiete wählen', '',
            'Geodaten (*.gpkg *.shp *.geojson *.json *.gz);;Alle Dateien (*)')
        if not path:
            return
        layer = None
        if path.lower().endswith('.gpkg'):
            subs = self._gpkg_sublayers(path)
            if len(subs) > 1:
                layer, ok = QInputDialog.getItem(self, 'Layer wählen', 'Layer mit den PLZ-Gebieten:',
                                                 subs, 0, False)
                if not ok:
                    return
            elif subs:
                layer = subs[0]
        lpath = ('/vsigzip/' + path) if path.lower().endswith('.gz') else path
        fields = self._layer_field_names(lpath, layer)
        if not fields:
            QMessageBox.warning(self, 'Datei nicht lesbar',
                                'Die Datei enthält keine lesbaren Flächen mit Attributen.')
            return
        pre = next((i for i, f in enumerate(fields)
                    if f.lower() in ('plz', 'postleitzahl', 'postcode', 'postal_code', 'zip')), 0)
        field, ok = QInputDialog.getItem(self, 'PLZ-Spalte', 'Spalte mit der Postleitzahl:',
                                         fields, pre, False)
        if not ok:
            return
        test = _plz.make_resolver(path, plz_field=field, layer_name=layer)
        if test is None or not test.ok:
            QMessageBox.warning(self, 'PLZ-Datei nicht verwendbar',
                                (test.error if test else '') or 'Die Datei kann nicht verwendet werden.')
            return
        cfg = {'path': path, 'field': field, 'layer': layer, 'name': os.path.basename(path)}
        _plz.set_override(cfg)
        import json as _json
        QgsSettings().setValue(PLZ_SETTINGS_KEY, _json.dumps(cfg))
        self._refresh_source_views()
        QMessageBox.information(
            self, 'PLZ-Gebiete',
            f'Eigene PLZ-Gebiete gesetzt ({_fmt_int(test.area_count)} Flächen). Die ergänzten PLZ '
            'werden beim nächsten „Alle eingetragenen Quellen aktualisieren“ neu berechnet.')

    def _use_bundled_plz(self):
        _plz.set_override(None)
        QgsSettings().remove(PLZ_SETTINGS_KEY)
        self._refresh_source_views()

    def _selected_source_bl(self):
        rows = self.tbl_sources.selectionModel().selectedRows() if self.tbl_sources.selectionModel() else []
        if not rows:
            return None
        it = self.tbl_sources.item(rows[0].row(), 0)
        return it.text() if it else None

    def _on_edit_button(self):
        bl = self.sender().property('bundesland') if self.sender() is not None else None
        if bl:
            self._open_source_detail(bl)

    def _on_source_double_clicked(self, item):
        it = self.tbl_sources.item(item.row(), 0)
        if it:
            self._open_source_detail(it.text())

    def _open_source_detail(self, bl):
        """Detail page of a Bundesland; afterwards starts the download/import if requested."""
        cat = self._catalog()
        entry = cat['sources'].get(bl) or {}
        try:
            standard = _sd.load_catalog(include_user=False)['sources'].get(bl) or {}
        except Exception:
            standard = {}
        running = self._import_running()
        meta = {} if running else (all_source_meta().get(bl) or {})
        adressen = (self._last_db_stats or {}).get(bl)
        _apply_proxy()                     # for „Quelle prüfen“
        dlg = SourceEditDialog(self, bl, entry, plugin_version=PLUGIN_VERSION, standard=standard,
                               meta=meta, adressen=adressen, busy=running or self._geocoding_running())
        dlg.exec()
        action, force = dlg.action, dlg.force
        dlg.deleteLater()
        self._refresh_source_views()
        if action == 'update':
            self._start_source_update([bl], force=force)

    def _auto_update_all(self):
        self._start_source_update(None, force=self.cb_auto_force.isChecked())

    def _geocoding_running(self):
        return self._thread is not None and self._thread.isRunning()

    def _import_running(self):
        return self._imp_thread is not None and self._imp_thread.isRunning()

    def _set_import_running(self, running):
        self._update_busy_state(import_running=running)

    def _update_busy_state(self, import_running=None, geocoding_running=None):
        """Geocoding and an import exclude each other (an import holds the
        database's write lock); the database path and the postcode areas cannot
        be changed while either runs."""
        imp = self._import_running() if import_running is None else import_running
        geo = self._geocoding_running() if geocoding_running is None else geocoding_running
        self.btn_run.setEnabled(not imp and not geo)
        self.btn_stop.setEnabled(geo and not self._cancel_requested)
        self.btn_import_stop.setEnabled(imp)
        for name in ('btn_db_save', 'btn_db_browse', 'btn_plz_own', 'btn_plz_bundled'):
            b = getattr(self, name, None)
            if b is not None:
                b.setEnabled(not imp and not geo)
        if getattr(self, 'btn_auto_update', None) is not None:
            self.btn_auto_update.setEnabled(not imp and not geo)
            self.btn_catalog.setEnabled(not imp and self._catalog_job is None)
        if not imp and not geo and getattr(self, 'btn_plz_bundled', None) is not None:
            self._update_plz_source_info()

    def _check_catalog(self):
        """„Quellenliste prüfen“ (check source list): load the online list in the background."""
        if self._catalog_job is not None:
            return
        self.btn_catalog.setEnabled(False)
        self.btn_catalog.setText('Prüfe …')
        _apply_proxy()
        thread = QThread()                 # no parent: may outlive the dialog (see shutdown)
        worker = CatalogCheckWorker()
        worker.moveToThread(thread)
        thread.started.connect(worker.run)
        worker.finished.connect(self._on_catalog_checked)   # method → runs in the UI thread
        worker.finished.connect(thread.quit)
        self._catalog_job = (thread, worker)                 # keep references
        thread.start()

    def _on_catalog_checked(self, status, text):
        job, self._catalog_job = self._catalog_job, None
        if job is not None:
            job[0].wait(2000)
        self.btn_catalog.setText('Quellenliste prüfen')
        self._tech_log(f'Quellenliste: {status} – {text}')
        self._refresh_source_views()
        if status in ('zu_neu', 'fehler'):
            QMessageBox.warning(self, 'Quellenliste', text)
        else:
            QMessageBox.information(self, 'Quellenliste', text)

    def _start_source_update(self, bundeslaender, force=False):
        """Download + import for the given Bundesländer (None = all with a
        registered source) in the background; progress at the bottom of the page."""
        if self._import_running():
            return
        if self._geocoding_running():
            QMessageBox.information(self, 'Geocodierung läuft',
                                    'Bitte warten, bis die Geocodierung abgeschlossen ist – während '
                                    'eines Imports ist die Datenbank gesperrt.')
            return
        cat = self._catalog()
        targets = bundeslaender or _sd.registered_bundeslaender(cat)
        if not targets:
            QMessageBox.information(self, 'Keine Quelle eingetragen',
                                    'Für kein Bundesland ist eine Quelle eingetragen.')
            return
        self.tab_widget.setCurrentWidget(self.tab_admin)
        self.admin_scroll.ensureWidgetVisible(self.grp_import_progress)
        _apply_proxy()
        self._set_import_running(True)
        self.import_log.clear()
        self.import_progress_bar.setMaximum(0)
        self.import_progress_bar.setValue(0)
        self._import_done = False
        if len(targets) == 1 and _sd.is_file_source(cat['sources'].get(targets[0])):
            self._import_log(f'Importiere Datei für {targets[0]} …')
        elif force:
            self._import_log('Lade neu (auch ohne neue Daten): ' + ', '.join(targets))
        else:
            self._import_log('Prüfe auf neue Daten: ' + ', '.join(targets))
        exp = [bl for bl in targets if (cat['sources'].get(bl) or {}).get('status') == 'experimentell']
        if exp:
            self._import_log('Hinweis: experimentelle Quelle(n): ' + ', '.join(exp)
                             + ' — Stand und Ergebnis bitte im Protokoll prüfen.')

        self._imp_thread = QThread()
        self._imp_worker = SourceUpdateWorker(bundeslaender, force=force, replace=True,
                                              plugin_version=PLUGIN_VERSION,
                                              check_catalog=not bundeslaender)
        self._imp_worker.moveToThread(self._imp_thread)
        self._imp_thread.started.connect(self._imp_worker.run)
        self._imp_worker.progress.connect(self._on_import_progress)
        self._imp_worker.log.connect(self._import_log)
        self._imp_worker.detail.connect(self._tech_log)
        self._imp_worker.finished.connect(self._on_import_finished)
        self._imp_thread.start()

    @staticmethod
    def _gpkg_sublayers(path):
        names = []
        try:
            from qgis.core import QgsProviderRegistry
            for sub in QgsProviderRegistry.instance().querySublayers(path):
                n = sub.name()
                if n and n not in names:
                    names.append(n)
        except Exception:
            pass
        if not names:
            try:
                from qgis.core import QgsVectorLayer
                vl = QgsVectorLayer(path, 'tmp', 'ogr')
                for s in vl.dataProvider().subLayers():
                    parts = s.split('!!::!!')
                    if len(parts) > 1 and parts[1] not in names:
                        names.append(parts[1])
            except Exception:
                pass
        return names

    @staticmethod
    def _layer_field_names(path, layername=None):
        try:
            from qgis.core import QgsVectorLayer
            uri = path if not layername else f'{path}|layername={layername}'
            vl = QgsVectorLayer(uri, 'tmp', 'ogr')
            if vl.isValid():
                return [f.name() for f in vl.fields()]
        except Exception:
            pass
        return []

    def _import_cancel(self):
        if self._imp_worker:
            self._imp_worker.cancel()
        self._import_log('Abbruch angefordert …')
        self.btn_import_stop.setEnabled(False)

    def _on_import_progress(self, current, total, status):
        # (the worker scales byte counts below 2^31 — QProgressBar is 32-bit)
        if total > 0:
            self.import_progress_bar.setMaximum(total)
            # Clamp the value to the maximum (the GML estimate can be too low)
            self.import_progress_bar.setValue(min(current, total))
        else:
            # Unknown total → animated busy bar
            if self.import_progress_bar.maximum() != 0:
                self.import_progress_bar.setMaximum(0)
        self.lbl_import_progress.setText(status)

    def _on_import_finished(self, success, message):
        # Prevent multiple triggering (the signal can fire again during
        # thread cleanup)
        if getattr(self, '_import_done', False):
            return
        self._import_done = True

        self._imp_thread.quit()
        self._imp_thread.wait()
        self._set_import_running(False)
        # Set the bar to completed
        self.import_progress_bar.setMaximum(100)
        self.import_progress_bar.setValue(100 if success else 0)
        aborted = 'abgebrochen' in message.lower() and 'fehler' not in message.lower()
        self.lbl_import_progress.setText(
            'Abgeschlossen.' if success else ('Abgebrochen.' if aborted else 'Mit Fehlern beendet.'))
        self._import_log(message)
        if not success and not aborted:
            self._import_log(_TECH_HINT)
        self._refresh_db_status()
        self._refresh_source_views()

        # Release worker/thread cleanly so that old signal connections
        # do not affect a subsequent import
        try:
            self._imp_worker.deleteLater()
        except Exception:
            pass
        self._imp_worker = None

        if success:
            QMessageBox.information(self, 'Quellen aktualisiert', message)
        elif aborted:
            QMessageBox.information(self, 'Quellen-Aktualisierung', message)
        else:
            QMessageBox.warning(self, 'Quellen-Aktualisierung', message)

    def _tech_log(self, msg):
        """Technical details to the QGIS message log (for support)."""
        try:
            QgsMessageLog.logMessage(str(msg), LOG_TAG, _MSG_INFO)
        except Exception:
            pass

    def _import_log(self, msg):
        self.import_log.append(msg)
        self.import_log.ensureCursorVisible()

    # ───────────────────────────────────── Log ──

    def _log(self, msg):
        self.log.append(msg)
        self.log.ensureCursorVisible()
